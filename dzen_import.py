# -*- coding: utf-8 -*-
"""Импорт официальной выгрузки Дзен Студии (xlsx «Статистика публикаций»).

Лист «Статистика»: по каждой публикации — дата, тип (Ролик/Статья/Видео),
название, ссылка, показы, открытия, дочитывания, время просмотра,
комментарии, подписки, лайки.

Публикации попадают в «Контент» (карточки со ссылками и статистикой),
а дневные суммы — в метрики канала Дзен (impressions, opens, reads,
likes, comments, subscribed).
"""
import warnings
from datetime import date, datetime

from db import db, Channel, ContentItem, ContentStat, RunLog, Notification

DZEN_CHANNEL_NAME = "Дзен"
FORMAT_MAP = {"ролик": "video", "видео": "video", "статья": "article",
              "пост": "post", "короткая статья": "article"}


def _get_channel():
    ch = Channel.query.filter_by(name=DZEN_CHANNEL_NAME).first()
    if not ch:
        ch = Channel(platform="dzen", name=DZEN_CHANNEL_NAME,
                     url="https://dzen.ru/alexey_syrover")
        db.session.add(ch)
        db.session.commit()
    return ch


def _num(v):
    if v in (None, "", "-"):
        return None
    try:
        return float(str(v).replace(" ", "").replace("\xa0", ""))
    except ValueError:
        return None


def import_dzen_xlsx(fileobj):
    """Главный вход: xlsx из Дзен Студии. Возвращает сводку."""
    warnings.filterwarnings("ignore")
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(fileobj.read()), data_only=True, read_only=False)
    ws = None
    for name in wb.sheetnames:
        if "статистик" in name.lower():
            ws = wb[name]
            break
    if ws is None:
        raise ValueError("Не найден лист «Статистика» — это выгрузка Дзен Студии?")

    # Найти строку заголовков: содержит «Дата публикации» и «Показы»
    header_row, cols = None, {}
    for r in range(1, min(10, ws.max_row) + 1):
        vals = {str(ws.cell(row=r, column=c).value or "").strip().lower(): c
                for c in range(1, ws.max_column + 1)}
        if "дата публикации" in vals and "показы" in vals:
            header_row, cols = r, vals
            break
    if header_row is None:
        raise ValueError("Не найдена строка заголовков выгрузки Дзена")

    def col(*keys):
        for k in keys:
            for name, c in cols.items():
                if k in name:
                    return c
        return None

    c_date = col("дата публикации")
    c_type = col("тип")
    c_title = col("публикация")
    c_link = col("ссылка")
    c_impr = col("показы")
    c_open = col("открытия")
    c_read = col("дочитывания")
    c_comm = col("комментарии")
    c_subs = col("подписки")
    c_like = col("лайки")

    run = RunLog(kind="dzen_import", status="OK")
    db.session.add(run)
    db.session.commit()

    ch = _get_channel()
    from connectors import save_metric

    stats = {"publications": 0, "days": 0}
    daily = {}
    existing = {(ci.link, ci.published_at) for ci in
                ContentItem.query.filter_by(channel_id=ch.id).all()}

    for r in range(header_row + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=c_date).value
        if v is None or (isinstance(v, str) and (not v.strip() or "всего" in v.lower())):
            continue
        if isinstance(v, datetime):
            pub_dt = v
        elif isinstance(v, date):
            pub_dt = datetime(v.year, v.month, v.day)
        else:
            try:
                pub_dt = datetime.strptime(str(v)[:19], "%Y-%m-%d %H:%M:%S")
            except ValueError:
                continue

        ptype = str(ws.cell(row=r, column=c_type).value or "").strip().lower() if c_type else ""
        fmt = FORMAT_MAP.get(ptype, "article")
        title = str(ws.cell(row=r, column=c_title).value or "")[:500] if c_title else ""
        link = str(ws.cell(row=r, column=c_link).value or "")[:500] if c_link else ""

        impressions = _num(ws.cell(row=r, column=c_impr).value if c_impr else None)
        opens = _num(ws.cell(row=r, column=c_open).value if c_open else None)
        reads = _num(ws.cell(row=r, column=c_read).value if c_read else None)
        comments = _num(ws.cell(row=r, column=c_comm).value if c_comm else None)
        subs = _num(ws.cell(row=r, column=c_subs).value if c_subs else None)
        likes = _num(ws.cell(row=r, column=c_like).value if c_like else None)

        # Контент: reach=показы, views=дочитывания (как «просмотры» единицы контента)
        if (link, pub_dt) not in existing:
            ci = ContentItem(channel_id=ch.id, link=link or None, published_at=pub_dt,
                             format=fmt, title=title or "(без названия)")
            db.session.add(ci)
            db.session.flush()
            cs = ContentStat(content_id=ci.id, date=pub_dt.date(),
                             reach=impressions, views=reads, likes=likes,
                             comments=comments, subs=subs)
            db.session.add(cs)
            stats["publications"] += 1

        # Дневная агрегация канала
        d = pub_dt.date()
        agg = daily.setdefault(d, {"impressions": 0, "opens": 0, "reads": 0,
                                   "likes": 0, "comments": 0, "subscribed": 0})
        for key, val in [("impressions", impressions), ("opens", opens), ("reads", reads),
                         ("likes", likes), ("comments", comments), ("subscribed", subs)]:
            if val:
                agg[key] += val

    for d, agg in sorted(daily.items()):
        for metric, val in agg.items():
            if val:
                save_metric(run.id, ch.id, d, metric, round(val, 0), "dzen_xlsx")

    db.session.commit()
    stats["days"] = len(daily)
    msg = (f"Дзен: публикаций {stats['publications']} (новых), дней со статистикой {stats['days']}")
    run.message = msg
    db.session.add(Notification(level="ok", message=msg))
    db.session.commit()
    return stats
